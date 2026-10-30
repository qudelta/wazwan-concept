"""Builds the editable PowerPoint from native/extract.mjs output.

Text is real, editable text (Archivo / JetBrains Mono, embedded in the file); cards, chips, bars and
rules are native shapes; icons, logos, the photo and charts are pictures. Each card becomes a group, and
each click in builds.json becomes a top-level group (or shape) with a Float In / Fade entrance.

Usage: python3 native/build_pptx.py <extract_dir> <out.pptx>
"""
import json
import math
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape as _escape

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

DECK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(DECK))
sys.path.append(str(DECK.parent / "session-01-deck"))   # shared: static fonts and EOT wrapper
import pptx_fonts  # noqa: E402
from pptx_common import effect_for, note_for, timing, transition  # noqa: E402

def escape(text: str) -> str:
    return _escape(text, {'"': "&quot;"})


EMU = 6350                       # per CSS px: 1920 px = 13.333 in
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = f'xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"'
FONT_DIR = DECK.parent / "session-01-deck" / "assets" / "pptx-fonts"

ARCHIVO = {400: ("Archivo", 0), 500: ("Archivo Medium", 0), 600: ("Archivo SemiBold", 0), 700: ("Archivo", 1),
           800: ("Archivo ExtraBold", 0), 900: ("Archivo Black", 0)}
ARCHIVO_SC = {400: ("Archivo SemiCondensed", 0), 700: ("Archivo SemiCondensed", 1),
              800: ("Archivo SemiCondensed ExtraBold", 0), 900: ("Archivo SemiCondensed Black", 0)}
MONO = {400: ("JetBrains Mono", 0), 500: ("JetBrains Mono Medium", 0), 700: ("JetBrains Mono", 1)}
used_faces: set[tuple[str, int]] = set()


def typeface(st) -> tuple[str, int]:
    table = MONO if st["family"] == "JetBrains Mono" else ARCHIVO_SC if st["stretch"] < 95 else ARCHIVO
    w = min(table, key=lambda k: (abs(k - st["weight"]), -k))
    used_faces.add(table[w])
    return table[w]


def e(px: float) -> int:
    return int(round(px * EMU))


def clr(c) -> str:
    r, g, b, a = c
    hexv = "".join(f"{max(0, min(255, round(v))):02X}" for v in (r, g, b))
    alpha = f'<a:alpha val="{round(a * 100000)}"/>' if a < 0.999 else ""
    return f'<a:srgbClr val="{hexv}">{alpha}</a:srgbClr>'


def blend_white(c):
    r, g, b, a = c
    return [r * a + 255 * (1 - a), g * a + 255 * (1 - a), b * a + 255 * (1 - a), 1]


class Ids:
    def __init__(self, start=5000):
        self.n = start

    def __call__(self):
        self.n += 1
        return self.n


def xfrm(x, y, w, h, rot=0, flipH=False, flipV=False) -> str:
    attrs = (f' rot="{rot * 60000}"' if rot else "") + (' flipH="1"' if flipH else "") + (' flipV="1"' if flipV else "")
    return f'<a:xfrm{attrs}><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{max(e(w), 1)}" cy="{max(e(h), 1)}"/></a:xfrm>'


def geometry(x, y, w, h, rad):
    """Pick a preset that rounds exactly the corners CSS rounds. Returns (xfrm, prstGeom)."""
    tl, tr, br, bl = (r > 0.5 for r in rad)
    r = max(rad)
    m = min(w, h)
    adj = lambda v: min(50000, round(v / m * 100000)) if m > 0 else 0  # noqa: E731
    if not any((tl, tr, br, bl)):
        return xfrm(x, y, w, h), '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
    if all((tl, tr, br, bl)):
        return xfrm(x, y, w, h), f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val {adj(r)}"/></a:avLst></a:prstGeom>'
    two = f'<a:prstGeom prst="round2SameRect"><a:avLst><a:gd name="adj1" fmla="val {adj(r)}"/><a:gd name="adj2" fmla="val 0"/></a:avLst></a:prstGeom>'
    if tl and tr and not (br or bl):
        return xfrm(x, y, w, h), two
    if bl and br and not (tl or tr):
        return xfrm(x, y, w, h, flipV=True), two
    cx, cy = x + w / 2, y + h / 2
    if tl and bl and not (tr or br):   # rotate the "top rounded" preset 270 degrees: swap width and height
        return xfrm(cx - h / 2, cy - w / 2, h, w, rot=270), two
    if tr and br and not (tl or bl):
        return xfrm(cx - h / 2, cy - w / 2, h, w, rot=90), two
    one = f'<a:prstGeom prst="round1Rect"><a:avLst><a:gd name="adj" fmla="val {adj(r)}"/></a:avLst></a:prstGeom>'
    if sum((tl, tr, br, bl)) == 1:
        return xfrm(x, y, w, h, flipH=tl or bl, flipV=bl or br), one
    return xfrm(x, y, w, h), f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val {adj(r)}"/></a:avLst></a:prstGeom>'


def sp(ids, name, x, y, w, h, fill=None, line=None, rad=(0, 0, 0, 0), shadow=None):
    sid = ids()
    xf, geom = geometry(x, y, w, h, rad)
    fill_xml = f"<a:solidFill>{clr(fill)}</a:solidFill>" if fill else "<a:noFill/>"
    if line:
        lw, lstyle, lcolor = line
        dash = '<a:prstDash val="dash"/>' if lstyle == "dashed" else '<a:prstDash val="sysDot"/>' if lstyle == "dotted" else ""
        ln = f'<a:ln w="{e(lw)}"><a:solidFill>{clr(lcolor)}</a:solidFill>{dash}</a:ln>'
    else:
        ln = "<a:ln><a:noFill/></a:ln>"
    eff = ""
    if shadow:
        dist = math.hypot(shadow["x"], shadow["y"])
        ang = (math.degrees(math.atan2(shadow["y"], shadow["x"])) % 360) if dist else 90
        eff = (f'<a:effectLst><a:outerShdw blurRad="{e(shadow["blur"])}" dist="{e(dist)}" dir="{round(ang * 60000)}" '
               f'algn="ctr" rotWithShape="0">{clr(shadow["color"])}</a:outerShdw></a:effectLst>')
    xml = (f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{sid}" name="{escape(name)} {sid}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
           f'<p:spPr>{xf}{geom}{fill_xml}{ln}{eff}</p:spPr>'
           f'<p:txBody><a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')
    return etree.fromstring(xml), (x, y, w, h), sid


def line_shape(ids, x1, y1, x2, y2, w, style, color):
    sid = ids()
    x, y, cw, ch = min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1)
    dash = '<a:prstDash val="dash"/>' if style == "dashed" else '<a:prstDash val="sysDot"/>' if style == "dotted" else ""
    xml = (f'<p:cxnSp {NS}><p:nvCxnSpPr><p:cNvPr id="{sid}" name="Rule {sid}"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
           f'<p:spPr>{xfrm(x, y, cw, ch)}<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
           f'<a:ln w="{e(w)}"><a:solidFill>{clr(color)}</a:solidFill>{dash}</a:ln></p:spPr></p:cxnSp>')
    return etree.fromstring(xml), (x, y, max(cw, 0.1), max(ch, 0.1)), sid


def shape_items(ids, it):
    """A CSS box: fill + uniform border as one shape; single-sided borders as thin bars or rules."""
    out = []
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    c = it["clip"]
    if c["w"] < w - 0.5 or c["h"] < h - 0.5:   # clipped by an overflow:hidden parent
        x, y, w, h = c["x"], c["y"], c["w"], c["h"]
    sides = it["sides"]
    on = [s for s in sides if s]
    uniform = len(on) == 4 and all(s["w"] == on[0]["w"] and s["style"] == on[0]["style"] and s["color"] == on[0]["color"] for s in on)
    if it["fill"] or uniform:
        line = None
        bx, by, bw, bh = x, y, w, h
        if uniform:
            lw = on[0]["w"]
            line = (lw, on[0]["style"], on[0]["color"])
            bx, by, bw, bh = x + lw / 2, y + lw / 2, w - lw, h - lw   # CSS borders sit inside the box
        out.append(sp(ids, "Box", bx, by, bw, bh, fill=it["fill"], line=line, rad=it["radius"], shadow=it["shadow"]))
    if not uniform:
        for k, s in enumerate(sides):
            if not s:
                continue
            bw = s["w"]
            if s["style"] in ("dashed", "dotted"):
                seg = [(x, y + bw / 2, x + w, y + bw / 2), (x + w - bw / 2, y, x + w - bw / 2, y + h),
                       (x, y + h - bw / 2, x + w, y + h - bw / 2), (x + bw / 2, y, x + bw / 2, y + h)][k]
                out.append(line_shape(ids, *seg, bw, s["style"], s["color"]))
            else:
                bx, by, bww, bhh = [(x, y, w, bw), (x + w - bw, y, bw, h), (x, y + h - bw, w, bw), (x, y, bw, h)][k]
                out.append(sp(ids, "Rule", bx, by, bww, bhh, fill=s["color"]))
    return out


def flatten(c, under):
    """Semi-transparent text over a solid colour -> the opaque colour it shows as."""
    if c[3] >= 0.999 or not under:
        return c
    a = c[3]
    return [c[k] * a + under[k] * (1 - a) for k in range(3)] + [1]


def text_item(ids, it):
    lines = it["lines"]
    under = it.get("under")
    L = it["lineHeight"]
    s0 = lines[0]["size"]
    top = lines[0]["base"] - (L - 0.2 * s0)
    minx, maxx = min(l["x0"] for l in lines), max(l["x1"] for l in lines)
    slack = (maxx - minx) * 0.05 + 12
    w = maxx - minx + slack
    x = minx if it["align"] == "l" else (minx + maxx) / 2 - w / 2 if it["align"] == "ctr" else maxx - w
    h = L * (len(lines) + sum(l.get("blankBefore", 0) for l in lines)) + 0.25 * s0
    paras, first_rpr = [], None
    for i, ln in enumerate(lines):
        for rn in ln["runs"]:
            st = rn["st"]
            face, bold = typeface(st)
            size = st["size"] / 0.83 if st["valign"] != "baseline" else st["size"]  # PowerPoint shrinks sub/superscripts itself
            attrs = f'lang="en-US" sz="{max(100, round(size * 50))}" b="{bold}" i="{1 if st["italic"] else 0}" kern="100"'
            if st["spacing"]:
                attrs += f' spc="{round(st["spacing"] * 50)}"'
            if st["caps"]:
                attrs += ' cap="all"'
            if st["valign"] == "sub":
                attrs += ' baseline="-25000"'
            elif st["valign"] == "super":
                attrs += ' baseline="30000"'
            hl = ""
            if st["highlight"]:
                c = st["highlight"]
                c = c if c[3] >= 0.99 else [c[0], c[1], c[2], c[3] * 0.7]
                hl = f'<a:highlight>{clr(blend_white(c))}</a:highlight>'
            rpr = (f'<a:rPr {attrs} dirty="0"><a:solidFill>{clr(flatten(st["color"], under))}</a:solidFill>{hl}'
                   f'<a:latin typeface="{escape(face)}"/><a:cs typeface="{escape(face)}"/></a:rPr>')
            first_rpr = first_rpr or rpr
            paras.append(f'<a:r>{rpr}<a:t>{escape(rn["text"])}</a:t></a:r>')
        if i < len(lines) - 1:
            br = ln["runs"][-1]["st"]
            for _ in range(lines[i + 1].get("blankBefore", 0)):
                face, bold = typeface(br)
                paras.append(f'<a:br><a:rPr lang="en-US" sz="{max(100, round(br["size"] * 50))}" b="{bold}">'
                             f'<a:latin typeface="{escape(face)}"/></a:rPr></a:br>')
            face, bold = typeface(br)
            paras.append(f'<a:br><a:rPr lang="en-US" sz="{max(100, round(br["size"] * 50))}" b="{bold}">'
                         f'<a:latin typeface="{escape(face)}"/></a:rPr></a:br>')
    end = first_rpr.replace("<a:rPr", "<a:endParaRPr", 1).replace("</a:rPr>", "</a:endParaRPr>")
    sid = ids()
    preview = escape("".join(r["text"] for r in lines[0]["runs"])[:40])
    wrap = "none" if it["nowrap"] else "square"
    xml = (f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{sid}" name="Text {sid}: {preview}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
           f'<p:spPr>{xfrm(x, top, w, h)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
           f'<p:txBody><a:bodyPr wrap="{wrap}" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" rtlCol="0"><a:noAutofit/></a:bodyPr>'
           f'<a:lstStyle/><a:p><a:pPr algn="{it["align"]}"><a:lnSpc><a:spcPts val="{max(0, round(L * 50))}"/></a:lnSpc>'
           f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>'
           f'{"".join(paras)}{end}</a:p></p:txBody></p:sp>')
    return etree.fromstring(xml), (x, top, w, h), sid


def group(ids, name, children):
    """children: list of (element, box, id). Group transform equals child space (no scaling)."""
    x0 = min(b[0] for _, b, _ in children)
    y0 = min(b[1] for _, b, _ in children)
    x1 = max(b[0] + b[2] for _, b, _ in children)
    y1 = max(b[1] + b[3] for _, b, _ in children)
    gid = ids()
    g = etree.fromstring(
        f'<p:grpSp {NS}><p:nvGrpSpPr><p:cNvPr id="{gid}" name="{escape(name)} {gid}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        f'<p:grpSpPr><a:xfrm><a:off x="{e(x0)}" y="{e(y0)}"/><a:ext cx="{e(x1 - x0)}" cy="{e(y1 - y0)}"/>'
        f'<a:chOff x="{e(x0)}" y="{e(y0)}"/><a:chExt cx="{e(x1 - x0)}" cy="{e(y1 - y0)}"/></a:xfrm></p:grpSpPr></p:grpSp>')
    for el, _, _ in children:
        g.append(el)
    return g, (x0, y0, x1 - x0, y1 - y0), gid


def build_slide(prs, idx, data, img_dir):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tree = slide.shapes._spTree
    ids = Ids()
    if data["bg"]:
        bg = etree.fromstring(f'<p:bg {NS}><p:bgPr><a:solidFill>{clr(data["bg"])}</a:solidFill><a:effectLst/></p:bgPr></p:bg>')
        slide._element.cSld.insert(0, bg)

    # 1. turn every item into PowerPoint elements (pictures go through python-pptx for the image part)
    produced = []   # (item, [(el, box, id), ...])
    for it in data["items"]:
        if it["kind"] == "raster":
            pic = slide.shapes.add_picture(str(img_dir / it["file"]), Emu(e(it["x"])), Emu(e(it["y"])), Emu(e(it["w"])), Emu(e(it["h"])))
            el = pic._element
            pid = ids()
            el.find(f".//{{{P}}}cNvPr").set("id", str(pid))
            el.find(f".//{{{P}}}cNvPr").set("name", f"Image {pid}")
            el.find(f".//{{{P}}}cNvPr").set("descr", it.get("alt", "Image"))
            el.getparent().remove(el)
            produced.append((it, [(el, (it["x"], it["y"], it["w"], it["h"]), pid)]))
        elif it["kind"] == "shape":
            produced.append((it, shape_items(ids, it)))
        elif it["kind"] == "text":
            produced.append((it, [text_item(ids, it)]))

    # 2. nest into card groups and click groups, keeping paint order
    top: list = []          # entries: ("el", tuple) or ("step", key) / ("card", key)
    steps: dict = {}        # stepEl -> {"step": n, "entries": [...], "cards": {}}
    cards_top: dict = {}
    for it, parts in produced:
        if not parts:
            continue
        holder = top
        if it["stepEl"]:
            if it["stepEl"] not in steps:
                steps[it["stepEl"]] = {"step": it["step"], "entries": [], "cards": {}}
                top.append(("step", it["stepEl"]))
            holder = steps[it["stepEl"]]["entries"]
            card_map = steps[it["stepEl"]]["cards"]
        else:
            card_map = cards_top
        if it["card"]:
            if it["card"] not in card_map:
                card_map[it["card"]] = []
                holder.append(("card", it["card"]))
            card_map[it["card"]].extend(parts)
        else:
            holder.extend(("el", p) for p in parts)

    def realize(entries, card_map):
        out = []
        for kind, val in entries:
            if kind == "el":
                out.append(val)
            else:
                members = card_map[val]
                out.append(members[0] if len(members) == 1 else group(ids, "Card", members))
        return out

    clicks: dict[int, list] = {}
    text_targets = set()
    for kind, val in top:
        if kind == "el":
            tree.append(val[0])
        elif kind == "card":
            el, _, _ = realize([("card", val)], cards_top)[0]
            tree.append(el)
        else:
            s = steps[val]
            members = realize(s["entries"], s["cards"])
            if not members:
                continue
            el, box, sid = members[0] if len(members) == 1 else group(ids, f"Click {s['step']}", members)
            tree.append(el)
            if el.tag == f"{{{P}}}sp" and el.find(f".//{{{P}}}txBody") is not None:
                text_targets.add(sid)
            clicks.setdefault(s["step"], []).append(sid)

    slide._element.append(transition())
    if clicks:
        slide._element.append(timing([[(sid, effect_for(idx, k)) for sid in clicks[k]] for k in sorted(clicks)], text_targets))
    note = note_for(idx, len(clicks))
    if note:
        slide.notes_slide.notes_text_frame.text = note
    return len(clicks)


def embed_fonts(pptx_path: Path) -> None:
    """Embed the used typefaces (EOT in ppt/fonts/*.fntdata), so the deck looks right without installing fonts."""
    by_family: dict[str, set[int]] = {}
    for face, bold in used_faces:
        by_family.setdefault(face, set()).add(bold)
    tmp = pptx_path.with_suffix(".tmp")
    with zipfile.ZipFile(pptx_path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        rels = etree.fromstring(zin.read("ppt/_rels/presentation.xml.rels"))
        pres = etree.fromstring(zin.read("ppt/presentation.xml"))
        ct = etree.fromstring(zin.read("[Content_Types].xml"))
        rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        used_ids = {r.get("Id") for r in rels}
        font_list = etree.SubElement(pres, f"{{{P}}}embeddedFontLst")
        n = 0
        fonts_data = {}
        for family in sorted(by_family):
            ef = etree.SubElement(font_list, f"{{{P}}}embeddedFont")
            etree.SubElement(ef, f"{{{P}}}font", typeface=family, pitchFamily="2" if "Mono" not in family else "49", charset="0")
            for bold in sorted({0} | by_family[family]):
                style = "Bold" if bold else "Regular"
                ttf = FONT_DIR / pptx_fonts.file_name(family, style)
                if not ttf.exists():
                    continue
                n += 1
                rid = f"rIdFont{n}"
                assert rid not in used_ids
                part = f"fonts/font{n}.fntdata"
                fonts_data[f"ppt/{part}"] = pptx_fonts.eot(ttf)
                etree.SubElement(rels, f"{{{rel_ns}}}Relationship", Id=rid, Target=part,
                                 Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
                etree.SubElement(ef, f"{{{P}}}{'bold' if bold else 'regular'}", {f"{{{R}}}id": rid})
        # embeddedFontLst must come right after notesSz (schema order)
        pres.remove(font_list)
        notes_sz = pres.find(f"{{{P}}}notesSz")
        notes_sz.addnext(font_list)
        pres.set("embedTrueTypeFonts", "1")
        pres.set("saveSubsetFonts", "1")
        ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
        if not any(d.get("Extension") == "fntdata" for d in ct.findall(f"{{{ct_ns}}}Default")):
            etree.SubElement(ct, f"{{{ct_ns}}}Default", Extension="fntdata", ContentType="application/x-fontdata")
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/_rels/presentation.xml.rels":
                data = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif item.filename == "ppt/presentation.xml":
                data = etree.tostring(pres, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif item.filename == "[Content_Types].xml":
                data = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
        for name, data in fonts_data.items():
            zout.writestr(name, data)
    tmp.replace(pptx_path)
    print(f"embedded {n} font files: {', '.join(sorted(by_family))}")


def main(extract_dir: str, out_path: str) -> None:
    src = Path(extract_dir)
    data = json.loads((src / "deck.json").read_text())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    total = sum(build_slide(prs, i, s, src / "img") for i, s in enumerate(data, 1))
    for n, slide in enumerate(prs.slides, 1):
        ids = [el.get("id") for el in slide._element.iter(f"{{{P}}}cNvPr")]
        assert len(ids) == len(set(ids)), f"slide {n}: duplicate shape ids"
    out = Path(out_path)
    prs.save(out)
    embed_fonts(out)
    print(f"pptx: {len(data)} slides, {total} click animations -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
